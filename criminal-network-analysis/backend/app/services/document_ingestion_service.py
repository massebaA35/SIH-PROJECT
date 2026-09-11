"""
Document ingestion: turn an uploaded file into hashed, stored case Evidence
and plain text ready for the existing offline NLP extraction pipeline.

Chain of custody starts at intake, not at commit: the SHA-256 hash is
computed on the untouched original bytes before anything else happens, and
that hash is also the deduplication key (the same document reappearing under
a different case is a signal worth surfacing, not hiding, so dedup is
checked across all cases, not just the target one).

No Evidence row or stored file is ever created for a rejected upload
(wrong type, too large, duplicate, unreadable) -- storage and the database
record only happen after extraction has already succeeded, so there is
never an Evidence row pointing at a file that isn't actually there or isn't
actually readable.
"""
from __future__ import annotations

import hashlib
import os
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from sqlalchemy.orm import Session

from app.config import get_settings
from app.models.evidence import Evidence

MAX_FILE_SIZE = 2 * 1024 * 1024


class UnsupportedFileType(Exception):
    """The uploaded file's extension isn't one this pipeline can read."""


class FileTooLarge(Exception):
    """The uploaded file exceeds MAX_FILE_SIZE."""


class UnreadableDocument(Exception):
    """The file parsed but yielded no usable text, or didn't parse at all."""


class DuplicateEvidence(Exception):
    """A file with this exact SHA-256 hash is already recorded as evidence."""

    def __init__(self, evidence_id: str, case_id: str):
        super().__init__(f"Duplicate of existing evidence {evidence_id} (case {case_id})")
        self.evidence_id = evidence_id
        self.case_id = case_id


def _extract_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # pypdf raises several distinct error types for malformed PDFs
        raise UnreadableDocument("Could not parse this PDF file. It may be corrupted.") from exc


def _extract_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(BytesIO(data))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)
    except Exception as exc:  # python-docx raises several distinct error types for malformed files
        raise UnreadableDocument("Could not parse this Word document. It may be corrupted.") from exc


_EXTRACTORS = {".txt": _extract_txt, ".pdf": _extract_pdf, ".docx": _extract_docx}
ALLOWED_EXTENSIONS = tuple(_EXTRACTORS)


def _extension_of(filename: str) -> str:
    return Path(filename.lower()).suffix


def hash_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def find_existing_evidence_by_hash(db: Session, sha256_hash: str) -> Evidence | None:
    return db.query(Evidence).filter(Evidence.sha256_hash == sha256_hash).first()


def _safe_filename(filename: str) -> str:
    """Strip any directory components so a crafted filename can't escape the
    per-case upload directory."""
    return os.path.basename(filename).strip() or "upload"


def store_file(case_id: str, evidence_id: str, filename: str, data: bytes) -> str:
    case_dir = Path(get_settings().uploads_dir) / case_id
    case_dir.mkdir(parents=True, exist_ok=True)
    path = case_dir / f"{evidence_id}_{_safe_filename(filename)}"
    path.write_bytes(data)
    return str(path)


def ingest_document(db: Session, case_id: str, filename: str, data: bytes, uploaded_by: str) -> dict:
    """Validate, hash, dedup-check, extract, store, and record one uploaded
    document as case Evidence. Raises one of the exceptions above on any
    rejection; returns {evidence_id, sha256_hash, text, file_path} on
    success. Commits the new Evidence row itself."""
    extension = _extension_of(filename)
    if extension not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '{extension or filename}'. Upload a .txt, .pdf, or .docx document."
        )
    if len(data) > MAX_FILE_SIZE:
        raise FileTooLarge("File is too large. Maximum size is 2 MB.")

    sha256_hash = hash_bytes(data)
    existing = find_existing_evidence_by_hash(db, sha256_hash)
    if existing:
        raise DuplicateEvidence(existing.id, existing.case_id)

    text = _EXTRACTORS[extension](data)
    if not text or not text.strip():
        raise UnreadableDocument(
            "No extractable text was found in this file. Scanned or image-only documents "
            "aren't supported yet -- try pasting the text instead."
        )

    evidence_id = f"EVID-{uuid4().hex[:8].upper()}"
    file_path = store_file(case_id, evidence_id, filename, data)

    db.add(Evidence(
        id=evidence_id,
        case_id=case_id,
        evidence_type="document",
        description=_safe_filename(filename),
        sha256_hash=sha256_hash,
        source="Uploaded document",
        uploaded_by=uploaded_by,
        uploaded_at=datetime.now(timezone.utc),
        file_path=file_path,
    ))
    db.commit()

    return {"evidence_id": evidence_id, "sha256_hash": sha256_hash, "text": text, "file_path": file_path}
